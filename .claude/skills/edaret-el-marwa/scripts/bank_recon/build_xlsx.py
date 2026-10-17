from common import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
NAVY="15294D"; GREY="EEF1F6"
F=lambda **k: Font(name="IBM Plex Sans Arabic",**k)
hdr=F(bold=True,color="FFFFFF",size=11); nf=F(size=10); bf=F(bold=True,size=10)
fill=PatternFill("solid",fgColor=NAVY); zf=PatternFill("solid",fgColor=GREY)
red=PatternFill("solid",fgColor="FDEEEE")
thin=Side(style="thin",color="CDD4E0"); bd=Border(thin,thin,thin,thin)
Cc=Alignment("center","center",wrap_text=True); Rr=Alignment("right","center",wrap_text=True)
wb=Workbook(); wb.remove(wb.active)
def sheet(title,headers,rows,widths,numcols=(),hl=None):
    ws=wb.create_sheet(title); ws.sheet_view.rightToLeft=True
    for j,h in enumerate(headers,1):
        c=ws.cell(1,j,h); c.font=hdr; c.fill=fill; c.alignment=Cc; c.border=bd
    for i,row in enumerate(rows,2):
        for j,v in enumerate(row,1):
            c=ws.cell(i,j,v); c.font=nf; c.border=bd
            c.alignment=Cc if isinstance(v,(int,float)) else Rr
            if j in numcols and isinstance(v,(int,float)): c.number_format='#,##0.00;[Red]-#,##0.00'
            if hl and hl(row): c.fill=red
            elif i%2==0: c.fill=zf
    for j,w in enumerate(widths,1): ws.column_dimensions[chr(64+j) if j<27 else 'A'].width=w
    ws.auto_filter.ref=f"A1:{chr(64+len(headers))}{len(rows)+1}"; ws.freeze_panes="A2"
# 1 summary
rows=[]
for k in ORDER:
    r=C[k]; cd,gd,pct,ng=coverage(r)
    rows.append([NAMES[k],r['last_bank_date'],r['stale_days'],r['bank_close'],r['odoo_at_last'],r['final_diff'],r['B0'],activity(r),
                 len(real_diffs(r)),ng,round(pct*100,1),'، '.join(s[0] for s in status(r)),r['slx_matched']['n']])
sheet("الملخص",["الحساب","آخر كشف مرفوع","تأخّر (يوم)","رصيد البنك (الكشف)","رصيد أودو في نفس التاريخ","الفرق الكلي (أودو − البنك)","منه: فرق افتتاحي ثابت","منه: فروق حركات","عدد السطور غير المتطابقة","عدد فجوات الكشوف","تغطية الكشوف %","الحالة","قيود تدقيق (uid 62) متطابقة مع البنك"],
      rows,[16,14,10,18,20,20,18,16,14,12,12,24,16],numcols=(4,5,6,7,8),hl=lambda r:abs(r[7])>0.5)
# 2 real diffs
rows=[]
for k in ORDER:
    for d in real_diffs(C[k]): rows.append([NAMES[k],d['kind'],d['date'],d['amt'],d['desc'],d['ref'],d['note']])
sheet("الفروق الحقيقية (سطور)",["الحساب","النوع","التاريخ","المبلغ","الوصف","رقم القيد في أودو","ملاحظة"],rows,[16,30,12,16,60,18,44],numcols=(4,))
# 3 timing
rows=[]
for k in ORDER:
    for s in C[k]['matched_shifted']: rows.append([NAMES[k],'نفس المبلغ بتاريخ مختلف (≤5 أيام)',s['bank_date'],s['odoo_date'],s['amt'],s['desc'],s['move']])
    for g in C[k]['groups']:
        sg=g['single']; side='سطر بنك = عدة سطور أودو' if g['single_side']=='bank' else 'سطر أودو = عدة سطور بنك'
        rows.append([NAMES[k],side,sg['date'],'، '.join(m['date'] for m in g['multi']),sg['amt'],'+'.join(f"{m['amt']:,.2f}" for m in g['multi']),sg.get('move','')])
sheet("فروق التوقيت والتجميع",["الحساب","النوع","تاريخ (جهة المفرد)","تواريخ الجهة الأخرى","المبلغ","الوصف/التفصيل","رقم القيد"],rows,[16,30,16,24,16,56,18],numcols=(5,))
# 4 chains & gaps
rows=[]
for k in ORDER:
    r=C[k]; first=True
    for x in r['rows']:
        rows.append([NAMES[k],'فترة بكشف' if x['type']=='chain' else 'أيام بلا كشف',x['from'],x['to'],x.get('n',x.get('days')),x['bank_net'],x['odoo_net'],x['diff'],x['running']])
sheet("الفترات والفجوات",["الحساب","النوع","من","إلى","عدد الحركات/الأيام","صافي البنك","صافي أودو","الفرق (أودو − البنك)","الفرق التراكمي"],rows,[16,16,12,12,12,18,18,18,18],numcols=(6,7,8,9),hl=lambda r:abs(r[7])>0.005)
# 5 opening
rows=[[NAMES[k],C[k]['anchor'],C[k]['B0']] for k in ORDER]
sheet("الفرق الافتتاحي",["الحساب","بداية الدفاتر المعتمدة","الفرق الافتتاحي (أودو − البنك) عند أول كشف"],rows,[18,22,34],numcols=(3,))
# 6 statements inventory
inv=[]
for p in sorted(P,key=lambda p:(p['key'],p['first'] or '')):
    inv.append([NAMES.get(p['key'],p['key']),p['name'],p['first'],p['last'],p['open'],p['close'],len(p['tx'])])
sheet("الكشوف المُحلَّلة",["الحساب","اسم الملف","أول حركة","آخر حركة","الرصيد الافتتاحي","رصيد الإقفال","عدد الحركات"],inv,[16,44,12,12,18,18,12],numcols=(5,6))
wb.save('recon.xlsx'); print('xlsx ok',[w.title for w in wb.worksheets],len(inv))
